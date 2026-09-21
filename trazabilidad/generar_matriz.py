#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# @artifact TOOL-MATRIZ
# @tipo Script
# @nombre Generador de la matriz de trazabilidad (caso Dietas al día)
# @version 0.1.0
# @estado Implementado
# @autor Santiago Maya Horta
# @cierre 2026-09-21
# @relacionados DOC-VAL, TEST-LOGIC, TEST-MANUAL, COD-ABOUT
"""
Genera trazabilidad/Matriz_de_trazabilidad.xlsx leyendo el repositorio (no se llena a mano).
La matriz es del caso 3 (Dietas al día, carpeta caso3-dietas), el único de los tres casos que tiene código.

De dónde sale cada dato
  - Requisitos, prioridad y criterios de aceptación: tabla de la sección 2 de caso3-dietas/docs/Documento_de_validacion.md.
  - Código que implementa cada requisito: etiquetas de requisito en los comentarios de caso3-dietas/src (ver README).
  - Casos de prueba: caso3-dietas/tests/logic.test.js (automáticos; se ejecutan con `node --test`) y
    caso3-dietas/tests/escenarios_manuales.md (manuales; el resultado lo escribe quien los ejecuta).
  - Metadatos de cada artefacto (ID, versión, estado final, autor, fecha de cierre, relacionados):
    caso 3: encabezado de cada archivo (y se regenera caso3-dietas/METADATOS.md);
    casos 1 y 2: tabla de casoN-.../METADATOS.md.

Uso (desde la raíz del repositorio)
    pip install openpyxl
    python trazabilidad/generar_matriz.py
    python trazabilidad/generar_matriz.py --sin-pruebas      # no ejecuta node; las pruebas automáticas quedan «Pendiente»
"""
import argparse
import datetime as dt
import math
import re
import shutil
import subprocess
import sys
from collections import OrderedDict, defaultdict
from copy import copy
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment
    from openpyxl.styles.cell_style import StyleArray
    from openpyxl.utils import get_column_letter as col
except ImportError:  # pragma: no cover
    sys.exit('Falta openpyxl. Instálalo con: pip install openpyxl')

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parent.parent
CASO3 = 'caso3-dietas'
BASE = ROOT / CASO3
DOC_REQ = f'{CASO3}/docs/Documento_de_validacion.md'
ABOUT = f'{CASO3}/src/screens/About.jsx'
MANUALES = f'{CASO3}/tests/escenarios_manuales.md'
PRIORIDAD_MOSCOW = {'Alta': 'Must have', 'Media': 'Should have', 'Baja': 'Could have'}
META_TAGS = ('artifact', 'tipo', 'nombre', 'version', 'estado', 'autor', 'cierre', 'relacionados')
ALL_TAGS = META_TAGS + ('req', 'infra')
META_ETIQUETA = {'artifact': 'ID único', 'version': 'versión', 'estado': 'estado final', 'autor': 'autor o revisor',
                 'cierre': 'fecha de cierre', 'relacionados': 'artefactos relacionados'}
OBLIGATORIOS = ('artifact', 'version', 'estado', 'autor', 'cierre', 'relacionados')
TAG_RE = re.compile(r'(?:^|[\s*/#>!-])@(' + '|'.join(ALL_TAGS) + r')\b[ \t]*(.*?)[ \t]*(?:-->|\*/)?[ \t]*$')
REQ_RE = re.compile(r'^(RF-A\d+|RNF-\d+|RF\d+)$')
CA_RE = re.compile(r'^CA\d+$')
DECL_RE = re.compile(r'^(\s*)(?:export\s+)?(?:default\s+)?(?:async\s+)?(?:function\s+(\w+)|(?:const|let|var)\s+(\w+)\s*=)')
HOY = dt.date.today().isoformat()


# ----------------------------------------------------------------------------- utilidades
def nat(s):
    return [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', s)]


def clave_req(i):
    return (0 if i.startswith('RF') else 1, nat(i))


def leer(rel):
    return (ROOT / rel).read_text(encoding='utf-8')


def archivos(*patrones):
    out = []
    for p in patrones:
        out += [x for x in sorted(ROOT.glob(p)) if x.is_file() and 'node_modules' not in x.parts]
    return [x.relative_to(ROOT).as_posix() for x in out]


def git(*args):
    try:
        r = subprocess.run(['git', *args], cwd=ROOT, capture_output=True, text=True, encoding='utf-8', timeout=30)
        return r.stdout.strip() if r.returncode == 0 else ''
    except Exception:
        return ''


def fecha_actualizacion(rel, cierre):
    if git('status', '--porcelain', '--', rel):
        return HOY  # modificado o nuevo y aún sin commit
    f = git('log', '-1', '--format=%cs', '--', rel)
    if f:
        return f
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}', cierre or ''):
        return cierre
    try:
        return dt.date.fromtimestamp((ROOT / rel).stat().st_mtime).isoformat()
    except OSError:
        return cierre or ''


def lista(valor):
    return [x for x in re.split(r'[,\s]+', valor or '') if x]


def leer_tags(lineas):
    out = []
    for i, l in enumerate(lineas):
        m = TAG_RE.search(l)
        if m:
            out.append((m.group(1), m.group(2).strip(), i))
    return out


def resolver_simbolo(lineas, idx):
    """Nombre de la función o constante a la que pertenece la etiqueta de la línea idx."""
    top = ''
    for k in range(idx, -1, -1):
        m = DECL_RE.match(lineas[k])
        if m and m.group(1) == '':
            top = m.group(2) or m.group(3)
            break
    j = idx + 1
    while j < len(lineas) and (not lineas[j].strip() or lineas[j].lstrip().startswith(('//', '*', '/*'))):
        j += 1
    if j < len(lineas):
        m = DECL_RE.match(lineas[j])
        if m:
            nombre = m.group(2) or m.group(3)
            if m.group(1) == '':
                return nombre
            return f'{top} › {nombre}' if top else nombre
    return top


def etiqueta_codigo(rel, simbolo):
    for pref in (f'{CASO3}/src/', f'{CASO3}/'):
        if rel.startswith(pref):
            rel = rel[len(pref):]
            break
    corto = rel
    return f'{corto} › {simbolo}' if simbolo else corto


# ----------------------------------------------------------------------------- lectura del documento de requisitos
def seccion(md, patron):
    lineas = md.splitlines()
    ini = next((i for i, l in enumerate(lineas) if re.match(patron, l)), None)
    if ini is None:
        return []
    fin = next((i for i in range(ini + 1, len(lineas)) if re.match(r'^## ', lineas[i])), len(lineas))
    return lineas[ini + 1:fin]


def tabla_md(lineas):
    filas, enc = [], None
    for l in lineas:
        if l.strip().startswith('|'):
            celdas = [c.strip() for c in l.strip().strip('|').split('|')]
            if enc is None:
                enc = celdas
                continue
            if all(re.fullmatch(r':?-{3,}:?', c) for c in celdas):
                continue
            filas.append(dict(zip(enc, celdas)))
        elif enc is not None:
            break
    return filas


def recortar(t, n):
    t = (t or '').strip()
    return t if len(t) <= n else t[:n - 1].rstrip() + '…'


def resumir(t, n):
    """Primeras oraciones del texto sin pasar de n caracteres (corta en el límite de una oración)."""
    t = (t or '').strip()
    if len(t) <= n:
        return t
    out = ''
    for s in re.split(r'(?<=[.;])\s+', t):
        if out and len(out) + 1 + len(s) > n:
            break
        out = f'{out} {s}'.strip()
    return out if 0 < len(out) <= n else recortar(t, n)


def leer_requisitos():
    md = leer(DOC_REQ)
    reqs = OrderedDict()
    for f in tabla_md(seccion(md, r'^## 2\.')):
        i = f.get('ID', '').strip()
        if REQ_RE.match(i):
            reqs[i] = {'id': i, 'texto': f.get('Requisito', ''), 'ca': f.get('Criterio de aceptación', ''),
                       'prioridad': f.get('Prioridad', ''), 'justificacion': f.get('Justificación', ''),
                       'codigo': [], 'pruebas': []}
    cas = {}
    for l in seccion(md, r'^## 6\.'):
        m = re.match(r'^- \*\*(CA\d+)[^*]*\*\*\s*(.+)$', l.strip())
        if m:
            cas[m.group(1)] = m.group(2).strip()
    m = re.search(r'\*\*Historia épica (EPC\d+)\.\*\*\s*Como (.+?) quiero', md)
    epica = m.group(1) if m else 'EPC'
    stakeholder = (m.group(2)[0].upper() + m.group(2)[1:]) if m else ''
    return reqs, cas, epica, stakeholder


def leer_about():
    if not (ROOT / ABOUT).exists():
        return {}
    t = leer(ABOUT)
    rx = re.compile(r"\{\s*id:\s*'([^']+)',\s*p:\s*'([^']+)',\s*ca:\s*'([^']+)',\s*t:\s*'((?:[^'\\]|\\.)*)'")
    return {m.group(1): {'p': m.group(2), 'ca': m.group(3), 't': m.group(4)} for m in rx.finditer(t)}


# ----------------------------------------------------------------------------- lectura del repositorio
def escanear(reqs):
    """Devuelve artefactos, pruebas, referencias a requisitos y hallazgos de estructura."""
    artefactos = OrderedDict()      # id -> dict
    sin_meta = []                   # archivos sin @artifact
    refs = defaultdict(set)         # id de requisito citado -> {archivos}
    sin_req = []                    # código sin requisito ni @infra
    pruebas = OrderedDict()         # TC-id -> dict

    grupos = [('codigo', archivos(f'{CASO3}/src/**/*.js', f'{CASO3}/src/**/*.jsx', f'{CASO3}/public/sw.js')),
              ('prueba', archivos(f'{CASO3}/tests/*.test.js')),
              ('manual', [MANUALES] if (ROOT / MANUALES).exists() else []),
              ('doc', archivos(f'{CASO3}/README.md', f'{CASO3}/docs/*.md')),
              ('script', archivos('trazabilidad/*.py'))]
    for grupo, rutas in grupos:
        for rel in rutas:
            lineas = leer(rel).splitlines()
            tags = leer_tags(lineas)
            meta = {}
            for tag, val, _ in tags:
                if tag in META_TAGS and tag not in meta:
                    meta[tag] = val
            usa_req = grupo in ('codigo', 'prueba')
            infra = usa_req and any(t == 'infra' for t, _, _ in tags)
            art_id = meta.get('artifact') or f'SIN-ID:{rel}'
            if 'artifact' not in meta:
                sin_meta.append(rel)
            art = {'id': art_id, 'ruta': rel, 'grupo': grupo, 'meta': meta, 'reqs': set(), 'infra': infra}
            artefactos[art_id] = art

            if usa_req:
                for tag, val, idx in tags:
                    if tag != 'req':
                        continue
                    simbolo = resolver_simbolo(lineas, idx) if grupo == 'codigo' else ''
                    for tok in lista(val.replace(',', ' ')):
                        if CA_RE.match(tok):
                            continue
                        if not REQ_RE.match(tok):
                            continue
                        art['reqs'].add(tok)
                        refs[tok].add(rel)
                        if grupo == 'codigo' and tok in reqs:
                            par = (rel, simbolo)
                            if par not in reqs[tok]['codigo']:
                                reqs[tok]['codigo'].append(par)
                if grupo == 'codigo' and not art['reqs'] and not infra:
                    sin_req.append(rel)

            if grupo == 'prueba':
                rx = re.compile(r"""\btest\(\s*(['"`])(TC-\d+)\s*[:\-–]\s*(.+?)\1""")
                for i, l in enumerate(lineas):
                    m = rx.search(l)
                    if not m:
                        continue
                    ids, k = set(), i - 1
                    while k >= 0 and lineas[k].strip().startswith(('//', '*', '/*')):
                        for tag, val, _ in leer_tags([lineas[k]]):
                            if tag == 'req':
                                ids |= {t for t in lista(val.replace(',', ' ')) if REQ_RE.match(t)}
                        k -= 1
                    for t in ids:
                        refs[t].add(rel)
                    pruebas[m.group(2)] = {'id': m.group(2), 'titulo': m.group(3).strip(), 'tipo': 'Automática',
                                           'reqs': sorted(ids, key=clave_req), 'resultado': 'Pendiente', 'fecha': '',
                                           'contenedor': art_id, 'ruta': rel, 'obs': ''}
            if grupo == 'manual':
                for l in lineas:
                    if not l.strip().startswith('| TC-'):
                        continue
                    c = [x.strip() for x in l.strip().strip('|').split('|')]
                    if len(c) < 7:
                        continue
                    ids = [t for t in lista(c[1].replace(',', ' ')) if REQ_RE.match(t)]
                    for t in ids:
                        refs[t].add(rel)
                    res = c[4] if c[4] in ('Pasado', 'Fallido', 'Bloqueado', 'Pendiente') else 'Pendiente'
                    pruebas[c[0]] = {'id': c[0], 'titulo': c[2], 'tipo': 'Manual', 'reqs': sorted(ids, key=clave_req),
                                     'resultado': res, 'fecha': c[5], 'contenedor': art_id, 'ruta': rel, 'obs': c[6]}
    for md in sorted(ROOT.glob('caso*/METADATOS.md')):
        if md.parent == BASE:
            continue  # el del caso 3 se genera automáticamente
        for f in tabla_md(md.read_text(encoding='utf-8').splitlines()):
            aid = f.get('ID', '').strip()
            if not aid:
                continue
            ruta = f"{md.parent.name}/{f.get('Archivo', '').strip().strip('`')}"
            meta = {'artifact': aid, 'tipo': f.get('Tipo', ''), 'nombre': f.get('Artefacto', ''), 'version': f.get('Versión', ''),
                    'estado': f.get('Estado final', ''), 'autor': f.get('Autor o revisor', ''),
                    'cierre': f.get('Fecha de cierre', ''), 'relacionados': f.get('Artefactos relacionados', '')}
            artefactos[aid] = {'id': aid, 'ruta': ruta, 'grupo': 'externo', 'meta': meta, 'reqs': set(), 'infra': False,
                               'existe': (ROOT / ruta).is_file()}
    return artefactos, pruebas, refs, sin_meta, sin_req


def ejecutar_pruebas_node(pruebas):
    """Ejecuta las pruebas automáticas y devuelve un mensaje de estado."""
    if shutil.which('node') is None:
        return 'node no está instalado: las pruebas automáticas quedan en «Pendiente».'
    out = ''
    for args in (['node', '--test', '--test-reporter=tap'], ['node', '--test']):
        try:
            r = subprocess.run(args, cwd=BASE, capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=240)
        except Exception as e:  # pragma: no cover
            return f'No se pudieron ejecutar las pruebas ({e}).'
        if 'bad option' in (r.stderr or '').lower():
            continue
        out = r.stdout
        break
    n = 0
    for m in re.finditer(r'^(not )?ok \d+ - (TC-\d+)\b(.*)$', out, re.M):
        tc = m.group(2)
        if tc not in pruebas:
            continue
        if re.search(r'#\s*(SKIP|TODO)', m.group(3), re.I):
            continue
        pruebas[tc]['resultado'] = 'Fallido' if m.group(1) else 'Pasado'
        pruebas[tc]['fecha'] = HOY
        n += 1
    if n == 0:
        return 'No se pudo leer el resultado de node --test; las pruebas automáticas quedan en «Pendiente».'
    return f'{n} pruebas automáticas ejecutadas.'


# ----------------------------------------------------------------------------- modelo
def construir_modelo(reqs, pruebas, artefactos, about, cas, epica, stakeholder):
    for tc in pruebas.values():
        for r in tc['reqs']:
            if r in reqs and tc['id'] not in reqs[r]['pruebas']:
                reqs[r]['pruebas'].append(tc['id'])
    doc = next((a for a in artefactos.values() if a['ruta'] == DOC_REQ), None)
    doc_ref = f"{doc['id']} v{doc['meta'].get('version', '?')}" if doc else ''
    filas = []
    for i in sorted(reqs, key=clave_req):
        r = reqs[i]
        tcs = sorted(r['pruebas'], key=nat)
        res = [pruebas[t]['resultado'] for t in tcs]
        if not r['codigo']:
            estado = 'Propuesto'
        elif 'Fallido' in res:
            estado = 'En desarrollo'
        elif res and all(x == 'Pasado' for x in res):
            estado = 'Verificado'
        else:
            estado = 'Implementado'
        if 'Fallido' in res:
            resultado = 'Fallido'
        elif 'Bloqueado' in res:
            resultado = 'Bloqueado'
        elif res and all(x == 'Pasado' for x in res):
            resultado = 'Pasado'
        else:
            resultado = 'Pendiente'
        ejecutadas = [pruebas[t]['fecha'] for t in tcs if pruebas[t]['resultado'] != 'Pendiente' and pruebas[t]['fecha']]
        codigo = sorted(r['codigo'])
        autores = sorted({artefactos_por_ruta(artefactos, rel)['meta'].get('autor', '') for rel, _ in codigo} - {''})
        ca = r['ca']
        ca_txt = cas.get(ca, '')
        criterios = f'{ca}: {resumir(ca_txt, 200)}' if ca_txt else ca
        n_ok = sum(1 for x in res if x == 'Pasado')
        n_fail = sum(1 for x in res if x == 'Fallido')
        n_pend = len(res) - n_ok - n_fail
        nota = f'Pruebas: {n_ok} pasadas, {n_fail} fallidas, {n_pend} pendientes.' if res else 'Sin pruebas asociadas.'
        corto = about.get(i, {}).get('t') or recortar(r['texto'], 58)
        filas.append({
            'id': i, 'titulo': corto, 'tipo': 'No funcional' if i.startswith('RNF') else 'Funcional',
            'necesidad': epica, 'stakeholder': stakeholder,
            'moscow': PRIORIDAD_MOSCOW.get(r['prioridad'], 'Should have'), 'version_doc': doc_ref,
            'criterios': criterios, 'codigo': codigo, 'tcs': tcs, 'resultado': resultado,
            'fecha_prueba': max(ejecutadas) if ejecutadas else '', 'estado': estado,
            'autores': ', '.join(autores) or '—',
            'actualizacion': max([fecha_actualizacion(rel, '') for rel, _ in codigo] or [HOY]), 'nota': nota,
        })
    return filas


def artefactos_por_ruta(artefactos, rel):
    return next(a for a in artefactos.values() if a['ruta'] == rel)


# ----------------------------------------------------------------------------- auditoría
def hallazgos(filas, reqs, pruebas, artefactos, refs, sin_meta, sin_req, about, cas, epica):
    h = []

    def add(sev, tipo, ref, desc, accion):
        h.append((sev, tipo, ref, desc, accion))

    for f in filas:
        if not f['codigo']:
            add('Alta', 'Requisito sin código', f['id'], 'Ningún archivo de src/ lo implementa (no hay etiqueta de requisito).',
                'Implementarlo o etiquetar el código que ya lo cubre.')
        if not f['tcs']:
            add('Alta', 'Requisito sin caso de prueba', f['id'], 'Ninguna prueba lo verifica.', 'Crear un caso de prueba y etiquetarlo con el requisito.')
        elif f['resultado'] == 'Fallido':
            add('Alta', 'Prueba fallida', f['id'], f"Hay pruebas fallidas: {', '.join(t for t in f['tcs'] if pruebas[t]['resultado'] == 'Fallido')}.",
                'Corregir el defecto y volver a ejecutar.')
        elif all(pruebas[t]['resultado'] == 'Pendiente' for t in f['tcs']):
            add('Media', 'Requisito sin prueba ejecutada', f['id'], f"Sus pruebas ({', '.join(f['tcs'])}) siguen en «Pendiente».",
                'Ejecutarlas y registrar el resultado en tests/escenarios_manuales.md.')
    for a in artefactos.values():
        if a['grupo'] == 'manual':
            pend = [t['id'] for t in pruebas.values() if t['tipo'] == 'Manual' and t['resultado'] == 'Pendiente']
            if pend:
                add('Media', 'Escenarios manuales sin ejecutar', a['id'], f"Pendientes: {', '.join(pend)}.",
                    'Ejecutarlos y anotar resultado, fecha y observaciones.')
    huerfanos = defaultdict(list)
    for i, archs in sorted(refs.items(), key=lambda kv: clave_req(kv[0])):
        if i not in reqs:
            huerfanos[tuple(sorted(archs))].append(i)
    for archs, ids in huerfanos.items():
        ref = f'{ids[0]}–{ids[-1]}' if len(ids) > 2 else ', '.join(ids)
        add('Baja', 'ID de requisito no definido en el documento', ref,
            f"Citado en {', '.join(archs)}, pero la sección 2 de {DOC_REQ} no lo define.",
            'Definir el requisito en el documento (o quitar la etiqueta si está fuera de alcance).')
    for rel in sin_req:
        add('Baja', 'Código sin requisito', rel, 'Implementa comportamiento que ningún requisito de la sección 2 respalda.',
            'Asociarlo a un requisito, o marcarlo con la etiqueta de infraestructura si no aplica.')
    for rel in sin_meta:
        add('Alta', 'Archivo sin metadatos', rel, 'No tiene el bloque de metadatos (ID, versión, estado, autor, cierre, relacionados).',
            'Agregar el encabezado de metadatos.')
    ids_validos = set(artefactos) | set(reqs) | set(pruebas)
    for a in artefactos.values():
        if a['grupo'] == 'externo' and not a.get('existe', True):
            add('Alta', 'Archivo declarado no existe', a['id'], f"{a['ruta']} no está en el repositorio.",
                'Copiar el archivo a esa carpeta con ese nombre, o corregir la columna «Archivo» de METADATOS.md.')
        falta = []
        for t in OBLIGATORIOS:
            v = a['meta'].get(t, '').strip()
            if not v or v.lower().startswith('pendiente'):
                falta.append(META_ETIQUETA[t])
        if falta and a['id'] not in [f"SIN-ID:{r}" for r in sin_meta]:
            add('Alta', 'Metadatos incompletos', a['id'], f"Falta o está pendiente: {', '.join(falta)}.",
                'Completar el encabezado del archivo.')
        for rel_id in lista(a['meta'].get('relacionados', '').replace(',', ' ')):
            if rel_id.lower() != 'ninguno' and rel_id not in ids_validos:
                add('Media', 'Artefacto relacionado inexistente', a['id'], f'«{rel_id}» no existe en el repositorio.',
                    'Corregir el ID o crear el artefacto.')
    if about:
        for i, r in reqs.items():
            if i not in about:
                add('Baja', 'Requisito ausente en la aplicación', i, f'Está en el documento pero no en la pantalla «Objetivo» ({ABOUT}).',
                    'Agregarlo a la lista de requisitos de la pantalla.')
            else:
                dif = []
                if about[i]['p'] != r['prioridad']:
                    dif.append(f"prioridad: documento «{r['prioridad']}», aplicación «{about[i]['p']}»")
                if about[i]['ca'] != r['ca']:
                    dif.append(f"criterio: documento «{r['ca']}», aplicación «{about[i]['ca']}»")
                if dif:
                    add('Media', 'Inconsistencia documento ↔ aplicación', i, '; '.join(dif) + '.', 'Alinear ambos.')
        for i in about:
            if i not in reqs:
                add('Media', 'Requisito de la aplicación no definido en el documento', i, f'Aparece en {ABOUT} pero no en la sección 2.', 'Definirlo en el documento.')
    usados = {r['ca'] for r in reqs.values()}
    for c in sorted(cas, key=nat):
        if c not in usados:
            add('Media', 'Criterio de aceptación sin requisito', c, 'Está propuesto en la sección 6 del documento, pero ningún requisito de la sección 2 lo usa; por eso no tiene código ni pruebas.',
                'Crear el requisito que lo cubra.')
    if cas:
        add('Media', 'Cambio propuesto sin RFC', f'{min(cas, key=nat)}–{max(cas, key=nat)}',
            'La sección 6 propone reformular los criterios de aceptación, pero no hay una RFC que registre el cambio.',
            'Registrar la RFC (formulario del curso) y agregarla al registro.')
    orden = {'Alta': 0, 'Media': 1, 'Baja': 2}
    h.sort(key=lambda x: (orden[x[0]], x[1], nat(x[2])))
    return h


# ----------------------------------------------------------------------------- Excel
def snap(ws, coord):
    return copy(ws[coord]._style)


def aplicar(cell, estilo, valor=None, **alineacion):
    cell._style = copy(estilo)
    if valor is not None:
        cell.value = valor
    if alineacion:
        a = copy(cell.alignment)
        cell.alignment = Alignment(horizontal=alineacion.get('horizontal', a.horizontal), vertical=alineacion.get('vertical', a.vertical),
                                   wrap_text=alineacion.get('wrap_text', a.wrap_text), text_rotation=alineacion.get('rotation', a.text_rotation))


def altura(textos_y_anchos, minimo=24, linea=12.5, extra=6):
    lineas = 1
    for texto, ancho in textos_y_anchos:
        if texto is None:
            continue
        cpl = max(int(ancho * 1.15), 4)
        n = sum(max(1, math.ceil(len(seg) / cpl)) for seg in str(texto).split('\n'))
        lineas = max(lineas, n)
    return max(minimo, lineas * linea + extra)


def hoja(wb, fragmento):
    return next(ws for ws in wb.worksheets if fragmento.lower() in ws.title.lower())


def limpiar(ws, r1, r2, c1, c2):
    vacio = StyleArray()
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell._style = copy(vacio)


def escribir_matriz(wb, filas):
    ws = hoja(wb, 'Matriz principal')
    letras = [col(i) for i in range(1, 19)]
    par = {c: snap(ws, f'{c}4') for c in letras}
    impar = {c: snap(ws, f'{c}5') for c in letras}
    s_prio = {'Must have': snap(ws, 'F4'), 'Should have': snap(ws, 'F7'), 'Could have': snap(ws, 'F13')}
    s_prio["Won't have"] = s_prio['Could have']
    s_tipo = {'Funcional': snap(ws, 'C4'), 'Seguridad': snap(ws, 'C6'), 'Interfaz': snap(ws, 'C7'),
              'Rendimiento': snap(ws, 'C12'), 'Restricción': snap(ws, 'C14'), 'No funcional': snap(ws, 'C12')}
    s_res = {'Pasado': snap(ws, 'M4'), 'Fallido': snap(ws, 'M11'), 'Pendiente': snap(ws, 'M13'), 'Bloqueado': snap(ws, 'M14'), 'N/A': snap(ws, 'M13')}
    s_est = {'Verificado': snap(ws, 'O4'), 'Aprobado': snap(ws, 'O4'), 'En desarrollo': snap(ws, 'O11'), 'Implementado': snap(ws, 'O11'),
             'Propuesto': snap(ws, 'O13')}
    limpiar(ws, 4, 103, 1, 18)
    anchos = {'A': 12, 'B': 34, 'C': 13, 'D': 14, 'E': 22, 'F': 12, 'G': 13, 'H': 46, 'I': 34, 'J': 12, 'K': 20, 'L': 24, 'M': 12, 'N': 13, 'O': 15, 'P': 12, 'Q': 14, 'R': 34}
    for k, v in anchos.items():
        ws.column_dimensions[k].width = v
    for i, f in enumerate(filas):
        r = 4 + i
        base = par if i % 2 == 0 else impar
        vals = [f['id'], f['titulo'], f['tipo'], f['necesidad'], f['stakeholder'], f['moscow'], f['version_doc'], f['criterios'],
                '\n'.join(etiqueta_codigo(a, b) for a, b in f['codigo']), 'N/A', f['autores'], ', '.join(f['tcs']),
                f['resultado'], f['fecha_prueba'], f['estado'], '', f['actualizacion'], f['nota']]
        for j, v in enumerate(vals):
            c = letras[j]
            est = base[c]
            if c == 'C':
                est = s_tipo[f['tipo']]
            elif c == 'F':
                est = s_prio[f['moscow']]
            elif c == 'M':
                est = s_res[f['resultado']]
            elif c == 'O':
                est = s_est[f['estado']]
            aplicar(ws.cell(r, j + 1), est, v if v != '' else None)
        ws.row_dimensions[r].height = altura([(vals[j], anchos[letras[j]]) for j in (1, 4, 7, 8, 10, 11, 17)], minimo=40)
    return {'Alta': s_prio['Must have'], 'Media': s_est['En desarrollo'], 'Baja': s_est['Propuesto']}


def escribir_visual(wb, filas, epica, pruebas, artefactos):
    ws = hoja(wb, 'Trazabilidad visual')
    t_titulo, t_g_nec, t_g_tc, t_g_cod = snap(ws, 'A1'), snap(ws, 'C2'), snap(ws, 'G2'), snap(ws, 'L2')
    t_h_id, t_h_nec, t_h_tc, t_h_cod = snap(ws, 'A3'), snap(ws, 'C3'), snap(ws, 'G3'), snap(ws, 'L3')
    t_id, t_tit = snap(ws, 'A4'), snap(ws, 'B4')
    t_ok_nec, t_ok_tc, t_ok_cod = snap(ws, 'C4'), snap(ws, 'G4'), snap(ws, 'L4')
    t_no_nec, t_no_tc, t_no_cod = snap(ws, 'D4'), snap(ws, 'H4'), snap(ws, 'M4')
    t_cob_lbl, t_cob_nec, t_cob_tc, t_cob_cod = snap(ws, 'A17'), snap(ws, 'C17'), snap(ws, 'G17'), snap(ws, 'L17')
    alto_titulo = ws.row_dimensions[1].height or 36
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    ws.delete_rows(1, ws.max_row + 5)
    for k in list(ws.column_dimensions):
        ws.column_dimensions[k].width = 9

    tcs = sorted({t for f in filas for t in f['tcs']}, key=nat)
    archivos_cod = sorted({rel for f in filas for rel, _ in f['codigo']})
    columnas = [('nec', epica, epica)] + [('tc', t, t) for t in tcs] + [('cod', rel, Path(rel).name) for rel in archivos_cod]
    c0 = 3
    ultima = c0 + len(columnas) - 1
    c_cod, c_tc = ultima + 1, ultima + 2
    total = c_tc

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total)
    aplicar(ws.cell(1, 1), t_titulo, 'TRAZABILIDAD VISUAL — Requisitos × Artefactos')
    ws.row_dimensions[1].height = alto_titulo
    grupos = [('nec', 'Necesidades', t_g_nec), ('tc', 'Casos de prueba', t_g_tc), ('cod', 'Módulos código', t_g_cod)]
    for g, nombre, est in grupos:
        idx = [c0 + i for i, (gg, _, _) in enumerate(columnas) if gg == g]
        if not idx:
            continue
        for c in range(idx[0], idx[-1] + 1):
            aplicar(ws.cell(2, c), est)
        ws.merge_cells(start_row=2, start_column=idx[0], end_row=2, end_column=idx[-1])
        ws.cell(2, idx[0]).value = nombre
    ws.row_dimensions[2].height = 18
    aplicar(ws.cell(3, 1), t_h_id, 'ID Req.')
    aplicar(ws.cell(3, 2), t_h_id, 'Título (abrev.)')
    h_est = {'nec': t_h_nec, 'tc': t_h_tc, 'cod': t_h_cod}
    for i, (g, _, etiqueta) in enumerate(columnas):
        aplicar(ws.cell(3, c0 + i), h_est[g], etiqueta, rotation=90, wrap_text=True, horizontal='center', vertical='center')
        ws.column_dimensions[col(c0 + i)].width = 11 if g == 'nec' else 5.5
    aplicar(ws.cell(3, c_cod), t_h_id, 'Nº archivos de código', rotation=90, wrap_text=True)
    aplicar(ws.cell(3, c_tc), t_h_id, 'Nº casos de prueba', rotation=90, wrap_text=True)
    ws.column_dimensions[col(c_cod)].width = 7
    ws.column_dimensions[col(c_tc)].width = 7
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 50
    ws.row_dimensions[3].height = 118
    ok = {'nec': t_ok_nec, 'tc': t_ok_tc, 'cod': t_ok_cod}
    no = {'nec': t_no_nec, 'tc': t_no_tc, 'cod': t_no_cod}
    cod_ini = next((c0 + i for i, (g, _, _) in enumerate(columnas) if g == 'cod'), None)
    tc_ini = next((c0 + i for i, (g, _, _) in enumerate(columnas) if g == 'tc'), None)
    for n, f in enumerate(filas):
        r = 4 + n
        aplicar(ws.cell(r, 1), t_id, f['id'])
        aplicar(ws.cell(r, 2), t_tit, f['titulo'])
        rels = {rel for rel, _ in f['codigo']}
        for i, (g, clave, _) in enumerate(columnas):
            hit = (g == 'nec') or (g == 'tc' and clave in f['tcs']) or (g == 'cod' and clave in rels)
            aplicar(ws.cell(r, c0 + i), ok[g] if hit else no[g], '✓' if hit else None)
        r1 = 4
        rc = f'{col(cod_ini)}{r}:{col(ultima)}{r}' if cod_ini else None
        rt = f'{col(tc_ini)}{r}:{col(cod_ini - 1 if cod_ini else ultima)}{r}' if tc_ini else None
        aplicar(ws.cell(r, c_cod), t_id, f'=COUNTIF({rc},"✓")' if rc else 0)
        aplicar(ws.cell(r, c_tc), t_id, f'=COUNTIF({rt},"✓")' if rt else 0)
        ws.row_dimensions[r].height = 20
    ult = 3 + len(filas)
    fr = ult + 2
    ws.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=2)
    aplicar(ws.cell(fr, 1), t_cob_lbl, 'Cobertura total (%)')
    aplicar(ws.cell(fr, 2), t_cob_lbl)
    cob = {'nec': t_cob_nec, 'tc': t_cob_tc, 'cod': t_cob_cod}
    for i, (g, _, _) in enumerate(columnas):
        L = col(c0 + i)
        aplicar(ws.cell(fr, c0 + i), cob[g], f'=IFERROR(COUNTIF({L}4:{L}{ult},"✓")/COUNTA($A$4:$A${ult}),"")')
    ws.row_dimensions[fr].height = 20
    ws.cell(fr + 1, 1).value = 'Cobertura de una columna = % de requisitos que ese artefacto cubre. Una celda vacía en una fila significa que el requisito no está cubierto por ese artefacto.'
    ws.freeze_panes = 'C4'


def escribir_resumen(wb, filas, matriz_titulo):
    ws = hoja(wb, 'Resumen')
    m = f"'{matriz_titulo}'"
    ws['G5'] = f'=SUMPRODUCT(--((({m}!O4:O103="Bloqueado")+({m}!M4:M103="Fallido")+({m}!M4:M103="Bloqueado"))>0))'
    ws['I5'] = f'=SUMPRODUCT(({m}!A4:A103<>"")*({m}!L4:L103=""))'
    for rng in list(ws.merged_cells.ranges):
        if str(rng) == 'A18:F18':
            ws.unmerge_cells('A18:F18')
    modelo = wb[matriz_titulo]
    t_id, t_txt, t_ctr = snap(modelo, 'A4'), snap(modelo, 'B4'), snap(modelo, 'D4')
    t_ok = snap(ws, 'A18')
    limpiar(ws, 18, 40, 1, 10)
    sin = [f for f in filas if not f['tcs']]
    if not sin:
        ws.merge_cells('A18:F18')
        aplicar(ws['A18'], t_ok, '✓ Todos los requisitos tienen al menos un caso de prueba asignado.')
        ws['A18'].font = copy(ws['A18'].font)
        ws.row_dimensions[18].height = 22
    else:
        for i, f in enumerate(sin):
            r = 18 + i
            for j, v in enumerate([f['id'], f['titulo'], f['moscow'], f['estado'], f['autores'], 'N/A']):
                aplicar(ws.cell(r, j + 1), t_id if j == 0 else (t_txt if j == 1 else t_ctr), v)
            ws.row_dimensions[r].height = 28


def escribir_rfc(wb, cas):
    ws = hoja(wb, 'Registro de RFC')
    t = snap(hoja(wb, 'Instrucciones'), 'A5')
    for rng in list(ws.merged_cells.ranges):
        if str(rng) != 'A1:M1':
            ws.unmerge_cells(str(rng))
    limpiar(ws, 3, 40, 1, 13)
    ws.merge_cells('A3:M3')
    msg = ('El caso 3 (Dietas al día) no tiene RFC registradas. ')
    if cas:
        msg += (f"El documento de validación propone reformular los criterios de aceptación ({', '.join(sorted(cas, key=nat))}), "
                'pero ese cambio aún no tiene una solicitud de cambio. Cuando exista, regístrala aquí con las columnas de arriba.')
    for c in range(1, 14):
        aplicar(ws.cell(3, c), t)
    ws['A3'] = msg
    ws.row_dimensions[3].height = 48


def escribir_catalogo(wb, artefactos, pruebas, reqs):
    ws = hoja(wb, 'Catálogo')
    h = snap(ws, 'A2')
    est = {'id': snap(ws, 'A3'), 'centro': snap(ws, 'D3'), 'texto': snap(ws, 'C3')}
    tipos = {'Documento': snap(ws, 'B3'), 'Documentación': snap(ws, 'B22'), 'Código': snap(ws, 'B14'), 'Suite de pruebas': snap(ws, 'B6'),
             'Caso de prueba': snap(ws, 'B6'), 'Script': snap(ws, 'B19'), 'Diagrama': snap(ws, 'B19')}
    limpiar(ws, 3, 60, 1, 11)
    t_titulo = snap(ws, 'A1')
    for rng in list(ws.merged_cells.ranges):
        if str(rng) == 'A1:G1':
            ws.unmerge_cells('A1:G1')
    ws.merge_cells('A1:K1')
    for c in range(2, 12):
        aplicar(ws.cell(1, c), t_titulo)
    for c, t in enumerate(['Estado final', 'Fecha de cierre', 'Artefactos relacionados', 'Ruta'], start=8):
        aplicar(ws.cell(2, c), h, t)
    anchos = {'A': 17, 'B': 16, 'C': 40, 'D': 10, 'E': 20, 'F': 14, 'G': 22, 'H': 15, 'I': 14, 'J': 34, 'K': 34}
    for k, v in anchos.items():
        ws.column_dimensions[k].width = v
    ws.row_dimensions[2].height = 34
    orden = {'externo': 0, 'doc': 1, 'codigo': 2, 'prueba': 3, 'manual': 3, 'script': 4}
    todos = sorted(artefactos.values(), key=lambda a: (orden[a['grupo']], a['id'] if a['grupo'] != 'externo' else nat(a['id'])))
    filas = []
    for a in todos:
        m = a['meta']
        vinc = ', '.join(sorted(a['reqs'], key=clave_req)) or '—'
        if a['ruta'] == DOC_REQ:
            vinc = ', '.join(sorted(reqs, key=clave_req))
        filas.append([a['id'], m.get('tipo', '—'), m.get('nombre', a['ruta']), m.get('version', '—'), m.get('autor', '—'),
                      fecha_actualizacion(a['ruta'], m.get('cierre', '')), vinc, m.get('estado', '—'), m.get('cierre', '—'),
                      m.get('relacionados', '—'), a['ruta']])
        if a['grupo'] in ('prueba', 'manual'):
            for tc in sorted((p for p in pruebas.values() if p['contenedor'] == a['id']), key=lambda p: nat(p['id'])):
                filas.append([tc['id'], 'Caso de prueba', tc['titulo'], m.get('version', '—'), m.get('autor', '—'),
                              fecha_actualizacion(a['ruta'], m.get('cierre', '')), ', '.join(tc['reqs']) or '—', m.get('estado', '—'),
                              m.get('cierre', '—'), ', '.join(tc['reqs'] + [a['id']]), a['ruta']])
    for i, v in enumerate(filas):
        r = 3 + i
        for j, x in enumerate(v):
            if j == 0:
                e = est['id']
            elif j == 1:
                e = tipos.get(x, tipos['Documento'])
            elif j in (2, 6, 9, 10):
                e = est['texto']
            else:
                e = est['centro']
            aplicar(ws.cell(r, j + 1), e, x)
        ws.row_dimensions[r].height = altura([(v[2], 40), (v[6], 22), (v[9], 34), (v[10], 34)], minimo=22)
    ws.freeze_panes = 'A3'


def escribir_auditoria(wb, items, sev):
    m = hoja(wb, 'Matriz principal')
    rfc = hoja(wb, 'Registro de RFC')
    ws = wb.create_sheet('🔎 Auditoría', index=wb.worksheets.index(hoja(wb, 'Resumen')) + 1)
    ws.sheet_view.showGridLines = False
    t_titulo, t_head = snap(rfc, 'A1'), snap(rfc, 'A2')
    t_num, t_txt, t_ctr = snap(rfc, 'A3'), snap(m, 'B4'), snap(m, 'D4')
    anchos = {'A': 6, 'B': 11, 'C': 30, 'D': 26, 'E': 62, 'F': 46}
    for k, v in anchos.items():
        ws.column_dimensions[k].width = v
    ws.merge_cells('A1:F1')
    aplicar(ws['A1'], t_titulo, 'AUDITORÍA DE TRAZABILIDAD (Misión 10) — hallazgos detectados por el script')
    ws.row_dimensions[1].height = 36
    for c, t in enumerate(['#', 'Severidad', 'Tipo de hallazgo', 'Artefacto / ID', 'Descripción', 'Acción sugerida'], start=1):
        aplicar(ws.cell(2, c), t_head, t)
    ws.row_dimensions[2].height = 30
    for i, (s, tipo, ref, desc, acc) in enumerate(items):
        r = 3 + i
        for c, (v, e) in enumerate([(i + 1, t_num), (s, sev[s]), (tipo, t_txt), (ref, t_ctr), (desc, t_txt), (acc, t_txt)], start=1):
            aplicar(ws.cell(r, c), e, v)
        ws.row_dimensions[r].height = altura([(tipo, 30), (ref, 26), (desc, 62), (acc, 46)], minimo=30)
    if not items:
        ws.merge_cells('A3:F3')
        aplicar(ws['A3'], t_txt, '✓ Sin hallazgos: la cadena necesidad → requisito → código → prueba está completa.')
    ws.freeze_panes = 'A3'


def escribir_instrucciones(wb, estado_pruebas):
    ws = hoja(wb, 'Instrucciones')
    t_h, t_l = snap(ws, 'A24'), snap(ws, 'A25')
    commit = git('rev-parse', '--short', 'HEAD')
    sucio = ' (con cambios sin confirmar)' if git('status', '--porcelain') else ''
    lineas = [
        f'Generada el {HOY} con trazabilidad/generar_matriz.py sobre el commit {commit or "(sin git)"}{sucio}. {estado_pruebas} No edites las hojas a mano: cambia el repositorio y vuelve a ejecutar el script.',
        'Caso elegido: caso 3, Dietas al día (caso3-dietas), el único con código. Requisitos, prioridad y criterios de aceptación: tabla de la sección 2 (y criterios de la sección 6) de caso3-dietas/docs/Documento_de_validacion.md. Prioridad Alta = Must have; Media = Should have.',
        'Código: etiquetas de requisito en los comentarios de caso3-dietas/src. Pruebas automáticas: caso3-dietas/tests/logic.test.js (resultado de la ejecución). Pruebas manuales: caso3-dietas/tests/escenarios_manuales.md (el resultado lo registra quien las ejecuta).',
        'Estado actual (regla): sin código = Propuesto · con código y alguna prueba sin pasar = Implementado · todas sus pruebas en Pasado = Verificado · alguna prueba Fallida = En desarrollo.',
        'Metadatos de todos los artefactos de los 3 casos (ID único, versión, estado final, autor o revisor, fecha de cierre y artefactos relacionados): hoja «Catálogo de artefactos». Caso 3: leídos del encabezado de cada archivo; casos 1 y 2: de su METADATOS.md.',
        'Cambio respecto a la plantilla original: las fórmulas de cobertura dividían entre las celdas con ✓ (siempre daban 100 %), «Sin caso de prueba» no descontaba las filas vacías y «Bloqueados / Fallidos» ignoraba el resultado «Bloqueado» de las pruebas. Ya están corregidas.',
        'Hoja «Auditoría»: rupturas detectadas por el script (requisitos sin código o sin prueba, IDs huérfanos, metadatos incompletos, archivos que faltan, inconsistencias documento ↔ aplicación).',
    ]
    r0 = 33
    for r in range(r0, r0 + len(lineas) + 2):
        for c in range(1, 9):
            ws.cell(r, c).value = None
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=8)
    for c in range(1, 9):
        aplicar(ws.cell(r0, c), t_h)
    ws.cell(r0, 1).value = 'CÓMO SE GENERÓ ESTA MATRIZ (caso 3: Dietas al día)'
    ws.row_dimensions[r0].height = 22
    for i, txt in enumerate(lineas):
        r = r0 + 1 + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        for c in range(1, 9):
            aplicar(ws.cell(r, c), t_l)
        ws.cell(r, 1).value = txt
        ws.row_dimensions[r].height = altura([(txt, 130)], minimo=26)


def escribir_metadatos_md(artefactos):
    """Regenera caso3-dietas/METADATOS.md (catálogo legible en GitHub) desde los encabezados de cada archivo."""
    orden = {'doc': 0, 'codigo': 1, 'prueba': 2, 'manual': 2, 'script': 3}
    propios = sorted((a for a in artefactos.values() if a['grupo'] != 'externo'), key=lambda a: (orden[a['grupo']], a['id']))
    lineas = [
        '# Metadatos — Caso 3: Dietas al día',
        '',
        '> Generado automáticamente por `trazabilidad/generar_matriz.py` a partir del encabezado de cada archivo. '
        'No lo edites a mano: cambia el encabezado del archivo y vuelve a ejecutar el script.',
        '',
        '| ID | Artefacto | Tipo | Archivo | Versión | Estado final | Autor o revisor | Fecha de cierre | Artefactos relacionados |',
        '|---|---|---|---|---|---|---|---|---|',
    ]
    for a in propios:
        m = a['meta']
        ruta = a['ruta'][len(CASO3) + 1:] if a['ruta'].startswith(CASO3 + '/') else '../' + a['ruta']
        celdas = [a['id'], m.get('nombre', ''), m.get('tipo', ''), f'`{ruta}`', m.get('version', ''), m.get('estado', ''),
                  m.get('autor', ''), m.get('cierre', ''), m.get('relacionados', '')]
        lineas.append('| ' + ' | '.join(celdas) + ' |')
    (BASE / 'METADATOS.md').write_text('\n'.join(lineas) + '\n', encoding='utf-8', newline='\n')


def main():
    ap = argparse.ArgumentParser(description='Genera la matriz de trazabilidad desde el repositorio.')
    ap.add_argument('--plantilla', default=str(ROOT / 'trazabilidad' / 'plantilla_matriz.xlsx'))
    ap.add_argument('--salida', default=str(ROOT / 'trazabilidad' / 'Matriz_de_trazabilidad.xlsx'))
    ap.add_argument('--sin-pruebas', action='store_true', help='no ejecuta node --test')
    a = ap.parse_args()

    reqs, cas, epica, stakeholder = leer_requisitos()
    if not reqs:
        sys.exit(f'No se encontraron requisitos en la sección 2 de {DOC_REQ}.')
    about = leer_about()
    artefactos, pruebas, refs, sin_meta, sin_req = escanear(reqs)
    estado_pruebas = 'Pruebas automáticas no ejecutadas (--sin-pruebas).' if a.sin_pruebas else ejecutar_pruebas_node(pruebas)
    print(estado_pruebas)
    filas = construir_modelo(reqs, pruebas, artefactos, about, cas, epica, stakeholder)
    items = hallazgos(filas, reqs, pruebas, artefactos, refs, sin_meta, sin_req, about, cas, epica)

    wb = openpyxl.load_workbook(a.plantilla)
    estilos_sev = escribir_matriz(wb, filas)
    escribir_visual(wb, filas, epica, pruebas, artefactos)
    escribir_resumen(wb, filas, hoja(wb, 'Matriz principal').title)
    escribir_rfc(wb, cas)
    escribir_catalogo(wb, artefactos, pruebas, reqs)
    escribir_auditoria(wb, items, estilos_sev)
    escribir_instrucciones(wb, estado_pruebas)
    wb.calculation.fullCalcOnLoad = True
    escribir_metadatos_md(artefactos)
    salida = Path(a.salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)

    print(f'Matriz guardada en {salida.relative_to(ROOT) if salida.is_relative_to(ROOT) else salida}')
    print(f'  Requisitos: {len(filas)} · Pruebas: {len(pruebas)} · Artefactos: {len(artefactos)} · Hallazgos: {len(items)}')
    for f in filas:
        print(f"  {f['id']:7} {f['estado']:14} código: {len(f['codigo'])}  pruebas: {len(f['tcs'])}  resultado: {f['resultado']}")


if __name__ == '__main__':
    main()
